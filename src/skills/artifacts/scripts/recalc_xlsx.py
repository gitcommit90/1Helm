#!/usr/bin/env python3
"""Recalculate every formula in a workbook with headless LibreOffice and report formula errors.

openpyxl saves formulas with no cached results, so until a real engine runs, every formula
cell reads back empty (in previews, in pandas, in load_workbook(data_only=True)). This script:
  1. has headless LibreOffice Calc load the workbook (recalculating everything) and export a copy;
  2. copies only the calculated values (and chart caches) from that copy into the ORIGINAL
     package, so styles, drawings and parts Calc would rewrite stay exactly as written;
  3. replaces the file in place and prints JSON.

JSON: {"status": "success"|"errors_found", "total_formulas", "total_errors", "error_summary"}.
"errors_found" still exits 0: read the JSON, not the exit code. An "error" key (and exit 1)
means nothing was recalculated and the file was not touched.

usage: recalc_xlsx.py <file.xlsx> [timeout_seconds] [--force]
  --force  recalculate even when cells linked to another workbook would lose their values
"""

import contextlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

if __name__ == "__main__" and any(a in ("-h", "--help") for a in sys.argv[1:]):
    print(__doc__.strip()); sys.exit(0)
from xml.etree.ElementTree import ParseError
from xml.parsers.expat import ExpatError

from openpyxl import load_workbook

from xlsx_formula_cache import merge_formula_caches

NO_SOFFICE = "soffice not found on PATH; LibreOffice is needed to recalculate"
MAX_LOCATIONS = 100
ERROR_VALUES = ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]
# A reference into another workbook: [1]Sheet!A1 or '[1]My sheet'!A1
EXTERNAL_REF_RE = re.compile(r"""(?<![\w"\[])'?\[\d+\][^!"\[\]]*'?!""")

# Profile setting: always recalculate OOXML/ODF files on load, so a plain --convert-to
# exports freshly calculated values. (No macro: a macro run opens a document window, which
# crashes LibreOffice 7.3 headless in some containers.)
RECALC_ON_LOAD = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def soffice():
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if binary is None:
        raise FileNotFoundError(NO_SOFFICE)
    return binary


def soffice_env(extra=None):
    """LibreOffice refuses to start without a writable HOME."""
    env = dict(os.environ)
    if not env.get("HOME") or not os.access(env["HOME"], os.W_OK):
        env["HOME"] = tempfile.gettempdir()
    env.update(extra or {})
    return env


def make_profile(profile: Path):
    """A fresh LibreOffice profile per run (parallel runs never share one), set to recalc on load."""
    user = profile / "user"
    user.mkdir(parents=True, exist_ok=True)
    (user / "registrymodifications.xcu").write_text(RECALC_ON_LOAD)
    return profile.as_uri()


def external_links_at_risk(filename):
    """Formula cells that read another workbook and have already lost their cached value.

    That file is not on this computer, so the cached value is the only data there is;
    recalculating would turn each of these cells into an error.
    """
    try:
        with zipfile.ZipFile(filename) as z:
            if not any(n.startswith("xl/externalLinks/") for n in z.namelist()):
                return []
    except (zipfile.BadZipFile, OSError):
        return []
    with contextlib.ExitStack() as stack:
        formulas = load_workbook(filename, data_only=False)
        stack.callback(formulas.close)
        values = load_workbook(filename, data_only=True)
        stack.callback(values.close)
        ext_names = [n for n, dn in formulas.defined_names.items()
                     if isinstance(getattr(dn, "value", None), str) and EXTERNAL_REF_RE.search(dn.value)]
        name_re = re.compile(r"\b(" + "|".join(map(re.escape, ext_names)) + r")\b") if ext_names else None
        at_risk = []
        for sheet in formulas.sheetnames:
            ws = formulas[sheet]
            if not hasattr(ws, "iter_rows"):
                continue
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("=") and (EXTERNAL_REF_RE.search(v) or (name_re and name_re.search(v))):
                        if values[sheet][cell.coordinate].value is None:
                            at_risk.append(f"{sheet}!{cell.coordinate}")
        return at_risk


def recalc(filename, timeout=60, force=False):
    path = Path(filename)
    if not path.exists():
        return {"error": f"{filename} does not exist"}
    abs_path = str(path.resolve())
    if not os.access(abs_path, os.W_OK):
        return {"error": f"{filename} is not writable; recalculation rewrites the file in place"}
    if not force:
        try:
            at_risk = external_links_at_risk(filename)
        except Exception as e:  # noqa: BLE001
            return {"error": f"could not inspect {filename} for links to other workbooks: {e}"}
        if at_risk:
            return {
                "error": ("Refusing to recalculate: this workbook links to another workbook, and "
                          f"{len(at_risk)} linked cell(s) have already lost their cached value (an openpyxl save strips "
                          "them). Recalculating would turn them into errors. Copy those values from the original file "
                          "before saving, or pass --force to accept the loss. Charts and conditional formats can also "
                          "hold such links, so this list may be incomplete."),
                "external_link_cells": at_risk[:MAX_LOCATIONS],
                "external_link_cells_truncated": max(0, len(at_risk) - MAX_LOCATIONS),
            }
    try:
        with tempfile.TemporaryDirectory(prefix="recalc-lo-", ignore_cleanup_errors=True) as tmp:
            original = Path(abs_path).read_bytes()
            calculated = Path(tmp) / "calculated.xlsx"
            result = run_calc(abs_path, calculated, timeout, Path(tmp) / "profile")
            if "error" in result:
                return result
            merged = merge_formula_caches(original, calculated)
            replace_file(Path(abs_path), original, merged)
            return result
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, ParseError, ExpatError) as e:
        return {"error": f"workbook was not replaced: {e}"}


def replace_file(path, original, merged):
    """Write the merged package next to the file, then swap it in; refuse if the file changed meanwhile."""
    path = path.resolve()
    tmp = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as out:
            tmp = Path(out.name)
            os.fchmod(out.fileno(), path.stat().st_mode & 0o777)
            out.write(merged)
        if path.read_bytes() != original:
            raise ValueError("the file changed during recalculation; run again on the current file")
        os.replace(tmp, path)
    finally:
        if tmp is not None:
            tmp.unlink(missing_ok=True)


def run_calc(abs_path, calculated, timeout, profile):
    url = make_profile(profile)
    outdir = calculated.parent / "out"
    outdir.mkdir(exist_ok=True)
    try:
        cmd = [soffice(), "--headless", "--norestore", f"-env:UserInstallation={url}",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir), abs_path]
    except FileNotFoundError:
        return {"error": NO_SOFFICE}
    if shutil.which("timeout"):
        cmd = ["timeout", str(timeout)] + cmd
    exported = outdir / (Path(abs_path).stem + ".xlsx")
    too_slow = f"LibreOffice timed out after {timeout}s; formulas were NOT recalculated. Run again with a longer timeout."
    r = None
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout + 15, env=soffice_env())
    except subprocess.TimeoutExpired:
        pass  # it can hang after storing; the exported file is what counts
    except FileNotFoundError:
        return {"error": NO_SOFFICE}
    if not exported.is_file() or exported.stat().st_size == 0:
        if r is None or r.returncode == 124:
            return {"error": too_slow}
        detail = (r.stderr or r.stdout or "").strip()[-400:] or f"exit {r.returncode}"
        return {"error": f"LibreOffice did not export a recalculated copy, so nothing was recalculated: {detail}"}
    os.replace(exported, calculated)
    try:
        return scan_errors(calculated)
    except Exception as e:  # noqa: BLE001
        return {"error": str(e)}


def scan_errors(calculated):
    wb = load_workbook(calculated, data_only=True)
    found = {e: [] for e in ERROR_VALUES}
    total = 0
    for name in wb.sheetnames:
        ws = wb[name]
        if not hasattr(ws, "iter_rows"):
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for e in ERROR_VALUES:
                        if e in cell.value:
                            found[e].append(f"{name}!{cell.coordinate}")
                            total += 1
                            break
    wb.close()
    out = {"status": "success" if total == 0 else "errors_found", "total_errors": total, "error_summary": {}}
    for e, locs in found.items():
        if locs:
            entry = {"count": len(locs), "locations": locs[:MAX_LOCATIONS]}
            if len(locs) > MAX_LOCATIONS:
                entry["locations_truncated"] = len(locs) - MAX_LOCATIONS
            out["error_summary"][e] = entry
    wbf = load_workbook(calculated, data_only=False)
    out["total_formulas"] = sum(1 for n in wbf.sheetnames if hasattr(wbf[n], "iter_rows")
                                for row in wbf[n].iter_rows() for c in row
                                if isinstance(c.value, str) and c.value.startswith("="))
    wbf.close()
    return out


def main():
    force = "--force" in sys.argv[1:]
    args = [a for a in sys.argv[1:] if a != "--force"]
    if not args:
        print(__doc__)
        sys.exit(1)
    result = recalc(args[0], int(args[1]) if len(args) > 1 else 60, force=force)
    print(json.dumps(result, indent=2))
    sys.exit(1 if "error" in result else 0)


if __name__ == "__main__":
    main()
