#!/usr/bin/env python3
"""readback.py: print a finished artifact's text the way its reader gets it, and scan it for leftover scaffolding.

Use it for the prose read-back and the placeholder scan before sharing anything. It reads the deliverable itself,
never the generator, so what you check is what the person receives.

  .pdf                page by page (pdftotext -layout)
  .docx               body paragraphs and tables in order, then headers and footers
  .pptx               slide text and speaker notes (a deck exported from images has only notes: read its .src HTML too)
  .xlsx / .xlsm       sheet by sheet: cell values (cached results) and how many formulas each sheet holds
  .csv / .tsv         row and column counts, ragged rows, then the rows
  .html / .htm        visible text (scripts and styles dropped); a deck's index.html is split per section.slide
  .md / .txt / other  the file as is

Usage:
  python3 readback.py FILE [--out FULL_TEXT_FILE] [--max-chars N] [--allow PATTERN ...]
stdout shows a summary, every placeholder hit, then the text up to --max-chars (default 12000; command output is
cut at 20k characters, so write the full text with --out and read it with read_file in parts when it is longer).
Exit 0 = no placeholders found, 1 = placeholders found (fix them, don't ship), 2 = the file could not be read.
--allow takes a regex for a hit that is genuinely part of the content (a document about TODO lists, say).
"""
import argparse, csv, html, json, os, re, subprocess, sys, zipfile

# Leftover scaffolding that must never reach the person. Case-insensitive.
PATTERNS = [
    (r"\bTODO\b|\bFIXME\b|\bTBD\b|\bFILL IN\b", "unfinished marker"),
    (r"lorem ipsum|\blorem\b|dolor sit amet", "filler text"),
    (r"\bplaceholder\b", "the word placeholder"),
    (r"\[(insert|your|add|name|date|company|link|image|photo|todo)[^\]]{0,40}\]", "bracketed slot"),
    (r"\b[xX]{3,}\b", "xxx run"),
    (r"\{\{[^}]{0,60}\}\}|\$\{[^}]{0,60}\}", "template variable"),
    (r"<(slug|artifact-slug|title|name|id)>", "angle-bracket slot"),
    (r"(image|visual|photo|chart|logo) (goes|will go) here", "empty slot note"),
    (r"\bsample text\b|\bexample\.com\b", "sample content"),
    (r"#REF!|#NAME\?|#DIV/0!|#VALUE!|#N/A\b|#NUM!", "formula error shown as text"),
    (r"Error! (Reference source|Bookmark) not", "broken Word field"),
    (r"\bundefined\b|\bNaN\b|\[object Object\]", "code value leaked into text"),
]


def run(cmd):
    return subprocess.run(cmd, capture_output=True, text=True, timeout=120)


def read_pdf(path):
    info = run(["pdfinfo", path])
    if info.returncode:
        raise RuntimeError(f"pdfinfo failed: {info.stderr.strip()[:200]}")
    pages = int(re.search(r"^Pages:\s+(\d+)", info.stdout, re.M).group(1))
    parts = []
    for i in range(1, pages + 1):
        t = run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), path, "-"]).stdout
        parts.append((f"page {i}", t.rstrip()))
    return parts, {"pages": pages}


def read_docx(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = Document(path)
    lines = []
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            style = p.style.name if p.style is not None else ""
            mark = "# " if style.startswith("Heading") or style == "Title" else ""
            if p.text.strip():
                lines.append(mark + p.text)
        elif tag == "tbl":
            for row in Table(el, d).rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
    head = []
    for s in d.sections:
        for part, name in ((s.header, "header"), (s.footer, "footer")):
            t = "\n".join(p.text for p in part.paragraphs if p.text.strip())
            if t:
                head.append(f"[{name}] {t}")
    meta = {"title_property": d.core_properties.title or "", "paragraphs": len(d.paragraphs), "tables": len(d.tables)}
    return [("body", "\n".join(lines)), ("headers and footers", "\n".join(dict.fromkeys(head)))], meta


def read_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts, text_shapes = [], 0
    for i, slide in enumerate(prs.slides, 1):
        bits = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                bits.append(sh.text_frame.text); text_shapes += 1
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            bits.append("[notes] " + slide.notes_slide.notes_text_frame.text)
        parts.append((f"slide {i}", "\n".join(bits)))
    meta = {"slides": len(prs.slides), "title_property": prs.core_properties.title or "", "text_shapes": text_shapes}
    if not text_shapes:
        meta["note"] = "image-only slides: also run readback on the deck's assembled .src/index.html"
    return parts, meta


def read_xlsx(path):
    from openpyxl import load_workbook
    vals = load_workbook(path, data_only=True)
    forms = load_workbook(path)
    parts, meta = [], {"sheets": [], "title_property": vals.properties.title or ""}
    for ws in vals.worksheets:
        fws = forms[ws.title]
        n_f = sum(1 for row in fws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        empty_f = 0
        rows = []
        for row in ws.iter_rows():
            cells = []
            for c in row:
                v = c.value
                if v is None and isinstance(fws[c.coordinate].value, str) and fws[c.coordinate].value.startswith("="):
                    empty_f += 1
                cells.append("" if v is None else str(v))
            if any(cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        meta["sheets"].append({"name": ws.title, "rows": len(rows), "formulas": n_f, "formulas_without_value": empty_f})
        parts.append((f"sheet {ws.title}", "\n".join(rows)))
    if any(s["formulas_without_value"] for s in meta["sheets"]):
        meta["note"] = "some formulas have no cached value: run recalc_xlsx.py first, or the person sees empty cells"
    return parts, meta


def read_csv(path):
    raw = open(path, newline="", encoding="utf-8-sig").read()
    dialect = csv.excel_tab if path.lower().endswith(".tsv") else csv.excel
    rows = list(csv.reader(raw.splitlines(), dialect))
    widths = {len(r) for r in rows if r}
    meta = {"rows": len(rows), "columns": max(widths) if widths else 0, "ragged": len(widths) > 1}
    return [("rows", "\n".join(" | ".join(r) for r in rows))], meta


def html_text(src):
    src = re.sub(r"(?is)<(script|style|template|noscript)\b.*?</\1>", " ", src)
    src = re.sub(r"(?is)<!--.*?-->", " ", src)
    src = re.sub(r"(?i)<br\s*/?>|</(p|div|h[1-6]|li|tr|section|table|ul|ol)>", "\n", src)
    src = re.sub(r"(?i)</t[dh]>", " | ", src)
    t = html.unescape(re.sub(r"<[^>]+>", " ", src))
    return "\n".join(re.sub(r"[ \t]+", " ", l).strip() for l in t.splitlines() if l.strip())


def read_html(path):
    src = open(path, encoding="utf-8", errors="replace").read()
    title = re.search(r"(?is)<title>(.*?)</title>", src)
    slides = re.split(r'(?i)(?=<section\b[^>]*class="[^"]*\bslide\b)', src)
    meta = {"title": html.unescape(title.group(1).strip()) if title else ""}
    if len(slides) > 2:
        parts = [(f"slide {i}", html_text(s)) for i, s in enumerate(slides[1:], 1)]
        meta["slides"] = len(parts)
        return parts, meta
    return [("page", html_text(src))], meta


def read_plain(path):
    return [("file", open(path, encoding="utf-8", errors="replace").read())], {}


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("file")
    ap.add_argument("--out")
    ap.add_argument("--max-chars", type=int, default=12000)
    ap.add_argument("--allow", action="append", default=[])
    a = ap.parse_args()
    ext = os.path.splitext(a.file)[1].lower()
    readers = {".pdf": read_pdf, ".docx": read_docx, ".dotx": read_docx, ".pptx": read_pptx, ".xlsx": read_xlsx, ".xlsm": read_xlsx,
               ".csv": read_csv, ".tsv": read_csv, ".html": read_html, ".htm": read_html}
    try:
        if ext in (".docx", ".pptx", ".xlsx", ".xlsm", ".dotx") and not zipfile.is_zipfile(a.file):
            raise RuntimeError("not a valid Office file (not a zip archive)")
        parts, meta = readers.get(ext, read_plain)(a.file)
    except Exception as e:  # report and stop: an unreadable file can't be shipped
        print(json.dumps({"ok": False, "file": a.file, "error": f"{type(e).__name__}: {e}"}))
        sys.exit(2)

    allow = [re.compile(p, re.I) for p in a.allow]
    hits = []
    for label, text in parts:
        for pat, why in PATTERNS:
            for m in re.finditer(pat, text, re.I):
                s = m.group(0)
                if any(x.search(s) for x in allow):
                    continue
                line = text[text.rfind("\n", 0, m.start()) + 1: text.find("\n", m.end()) if text.find("\n", m.end()) > -1 else len(text)]
                hits.append({"where": label, "found": s, "why": why, "line": line.strip()[:160]})
    full = "\n\n".join(f"===== {label} =====\n{text}" for label, text in parts)
    if a.out:
        os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
        open(a.out, "w").write(full)
    summary = {"ok": not hits, "file": a.file, **meta, "chars": len(full), "placeholders": hits[:50]}
    if len(hits) > 50:
        summary["placeholders_truncated"] = len(hits) - 50
    print(json.dumps(summary, indent=1, ensure_ascii=False))
    print()
    print(full[: a.max_chars])
    if len(full) > a.max_chars:
        print(f"\n[... {len(full) - a.max_chars} more characters{' in ' + a.out if a.out else '; pass --out FILE to keep the full text'}]")
    sys.exit(0 if not hits else 1)


if __name__ == "__main__":
    main()
